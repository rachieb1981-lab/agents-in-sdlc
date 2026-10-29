import base64
import io
import importlib.util
import sys
from pathlib import Path
import pytest
from unittest.mock import MagicMock, patch

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from PIL import Image

# Import target module with hyphenated name
script_path = Path(__file__).parent / "convert-pptx.py"
spec = importlib.util.spec_from_file_location("convert_pptx", script_path)
convert_pptx = importlib.util.module_from_spec(spec)
sys.modules["convert_pptx"] = convert_pptx
spec.loader.exec_module(convert_pptx)


class TestRGBToHex:
    def test_rgb_to_hex_valid(self):
        rgb = RGBColor(255, 0, 0)
        assert convert_pptx.rgb_to_hex(rgb) == "#FF0000"

    def test_rgb_to_hex_none(self):
        assert convert_pptx.rgb_to_hex(None) is None

    def test_rgb_to_hex_invalid_type(self):
        class BadRGB:
            def __str__(self):
                raise ValueError("Error")
        assert convert_pptx.rgb_to_hex(BadRGB()) is None


class TestEnsurePPTX:
    def test_ensure_pptx_success(self):
        assert convert_pptx._ensure_pptx() is True

    def test_ensure_pptx_import_error(self):
        with patch.dict(sys.modules, {"pptx": None}):
            with pytest.raises(SystemExit) as exc_info:
                convert_pptx._ensure_pptx()
            assert exc_info.value.code == 1


class TestGetTextStyle:
    def test_get_text_style_full(self):
        run = MagicMock()
        run.font.bold = True
        run.font.italic = True
        run.font.underline = True
        run.font.size.pt = 14
        run.font.color.rgb = RGBColor(0, 128, 255)
        run.font.name = "Arial"

        style = convert_pptx.get_text_style(run)
        assert "font-weight:bold" in style
        assert "font-style:italic" in style
        assert "text-decoration:underline" in style
        assert "font-size:14pt" in style
        assert "color:#0080FF" in style
        assert "font-family:'Arial',sans-serif" in style

    def test_get_text_style_empty(self):
        run = MagicMock()
        run.font.bold = False
        run.font.italic = False
        run.font.underline = False
        run.font.size = None
        run.font.color.rgb = None
        run.font.name = None

        style = convert_pptx.get_text_style(run)
        assert style == ""

    def test_get_text_style_exception(self):
        run = MagicMock()
        type(run).font = pytest.raises(AttributeError)
        style = convert_pptx.get_text_style(run)
        assert style == ""


class TestGetAlignment:
    def test_alignment_center(self):
        paragraph = MagicMock()
        paragraph.alignment = PP_ALIGN.CENTER
        assert convert_pptx.get_alignment(paragraph) == "center"

    def test_alignment_right(self):
        paragraph = MagicMock()
        paragraph.alignment = PP_ALIGN.RIGHT
        assert convert_pptx.get_alignment(paragraph) == "right"

    def test_alignment_justify(self):
        paragraph = MagicMock()
        paragraph.alignment = PP_ALIGN.JUSTIFY
        assert convert_pptx.get_alignment(paragraph) == "justify"

    def test_alignment_left_or_default(self):
        paragraph = MagicMock()
        paragraph.alignment = PP_ALIGN.LEFT
        assert convert_pptx.get_alignment(paragraph) == "left"

    def test_alignment_exception(self):
        paragraph = MagicMock()
        type(paragraph).alignment = pytest.raises(AttributeError)
        assert convert_pptx.get_alignment(paragraph) == "left"


class TestGetShapePosition:
    def test_valid_position(self):
        shape = MagicMock()
        shape.left = 100
        shape.top = 200
        shape.width = 300
        shape.height = 400
        left, top, width, height = convert_pptx.get_shape_position(shape, 1000, 1000)
        assert left == 10.0
        assert top == 20.0
        assert width == 30.0
        assert height == 40.0

    def test_zero_or_none_position(self):
        shape = MagicMock()
        shape.left = None
        shape.top = None
        shape.width = None
        shape.height = None
        left, top, width, height = convert_pptx.get_shape_position(shape, 1000, 1000)
        assert (left, top, width, height) == (0, 0, 50, 30)

    def test_exception_fallback(self):
        shape = MagicMock()
        type(shape).left = pytest.raises(TypeError)
        left, top, width, height = convert_pptx.get_shape_position(shape, 1000, 1000)
        assert (left, top, width, height) == (5, 5, 90, 40)


class TestGetSlideBackground:
    def test_background_default(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        bg = convert_pptx.get_slide_background(slide, prs)
        assert bg == "background-color:#ffffff"


class TestGetShapeFill:
    def test_no_fill(self):
        shape = MagicMock()
        shape._element.find.return_value = None
        assert convert_pptx.get_shape_fill(shape) is None


class TestRenderParagraph:
    def test_render_paragraph_escaping_and_style(self):
        run = MagicMock()
        run.text = "<Hello & Welcome>"
        run.font.bold = True
        run.font.italic = False
        run.font.underline = False
        run.font.size = None
        run.font.color.rgb = None
        run.font.name = None

        paragraph = MagicMock()
        paragraph.alignment = PP_ALIGN.CENTER
        paragraph.runs = [run]

        html = convert_pptx.render_paragraph(paragraph)
        assert "text-align:center" in html
        assert "&lt;Hello &amp; Welcome&gt;" in html
        assert 'style="font-weight:bold"' in html

    def test_render_paragraph_empty(self):
        paragraph = MagicMock()
        paragraph.runs = []
        assert convert_pptx.render_paragraph(paragraph) == ""


class TestExtractImageDataAndCountImages:
    def test_extract_image_data_success(self):
        shape = MagicMock()
        shape.image.blob = b"fake_image_bytes"
        shape.image.content_type = "image/png"
        blob, content_type = convert_pptx.extract_image_data(shape)
        assert blob == b"fake_image_bytes"
        assert content_type == "image/png"

    def test_extract_image_data_no_image(self):
        shape = MagicMock(spec=[])
        blob, content_type = convert_pptx.extract_image_data(shape)
        assert blob is None
        assert content_type is None

    def test_count_images(self):
        prs = MagicMock()
        slide1 = MagicMock()
        shape1 = MagicMock()
        shape1.shape_type = 13
        shape1.image = MagicMock()

        shape2 = MagicMock()
        shape2.shape_type = 1
        del shape2.image

        slide1.shapes = [shape1, shape2]
        prs.slides = [slide1]

        assert convert_pptx.count_images(prs) == 1


class TestConvertIntegration:
    @pytest.fixture
    def sample_pptx(self, tmp_path):
        prs = Presentation()
        # Slide 1: Title and text
        slide1 = prs.slides.add_slide(prs.slide_layouts[6])
        tx_box = slide1.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        tf = tx_box.text_frame
        p = tf.paragraphs[0]
        p.text = "Sample Presentation Title"

        # Slide 2: Table and image
        slide2 = prs.slides.add_slide(prs.slide_layouts[6])

        # Table
        table_shape = slide2.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(2))
        table = table_shape.table
        table.cell(0, 0).text = "Header 1"
        table.cell(0, 1).text = "Header 2"
        table.cell(1, 0).text = "Data & Info"
        table.cell(1, 1).text = "<Cell>"

        # Image
        img = Image.new("RGB", (100, 100), color="red")
        img_byte_arr = io.BytesIO()
        img.save(img_byte_arr, format="PNG")
        img_bytes = img_byte_arr.getvalue()
        img_path = tmp_path / "test_img.png"
        img_path.write_bytes(img_bytes)

        slide2.shapes.add_picture(str(img_path), Inches(1), Inches(3), Inches(2), Inches(2))

        pptx_file = tmp_path / "test_presentation.pptx"
        prs.save(str(pptx_file))
        return pptx_file

    def test_convert_inline_assets(self, sample_pptx, tmp_path):
        out_html = tmp_path / "output.html"
        convert_pptx.convert(str(sample_pptx), output_path=str(out_html), external_assets=False)

        assert out_html.exists()
        content = out_html.read_text(encoding="utf-8")
        assert "<title>Sample Presentation Title</title>" in content
        assert "Header 1" in content
        assert "Data &amp; Info" in content
        assert "data:image/png;base64," in content

    def test_convert_external_assets(self, sample_pptx, tmp_path):
        out_html = tmp_path / "output.html"
        convert_pptx.convert(str(sample_pptx), output_path=str(out_html), external_assets=True)

        assert out_html.exists()
        content = out_html.read_text(encoding="utf-8")
        assert "assets/img-001.png" in content
        assert (tmp_path / "assets" / "img-001.png").exists()

    def test_convert_auto_external_assets_trigger(self, sample_pptx, tmp_path):
        out_html = tmp_path / "output_auto.html"
        with patch.object(convert_pptx, "count_images", return_value=60):
            convert_pptx.convert(str(sample_pptx), output_path=str(out_html), external_assets=None)
            assert (tmp_path / "assets").exists()

    def test_convert_large_file_warning(self, sample_pptx, tmp_path, capsys):
        out_html = tmp_path / "output_large.html"
        with patch("os.path.getsize", return_value=160 * 1024 * 1024):
            convert_pptx.convert(str(sample_pptx), output_path=str(out_html), external_assets=False)
            captured = capsys.readouterr()
            assert "WARNING: File is 160MB" in captured.out
